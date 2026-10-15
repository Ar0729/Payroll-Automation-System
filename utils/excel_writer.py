from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def format_sheet(ws):

    # Header formatting
    for cell in ws[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # Auto column width
    for column in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            40
        )


def write_excel(results, output_file):

    with __import__("pandas").ExcelWriter(
        output_file,
        engine="openpyxl"
    ) as writer:

        for sheet_name, df in results.items():

            if df is None:
                continue

            if df.empty:

                df = __import__("pandas").DataFrame({
                    "Message": [
                        "No records found"
                    ]
                })

            df.to_excel(
                writer,
                sheet_name=sheet_name[:31],
                index=False
            )

    # Apply formatting
    wb = load_workbook(output_file)

    for ws in wb.worksheets:

        format_sheet(ws)

        ws.freeze_panes = "A2"

        ws.auto_filter.ref = ws.dimensions

    wb.save(output_file)

    print(
        f"Excel report generated: {output_file}"
    )